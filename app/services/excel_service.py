import io
from datetime import date
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from app.models.inventory import Material
from app.models.finance import CashTransaction
from app.models.sales import Order
from app.config import FACTORY_NAME

def create_header_style():
    fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    return fill, font, align, border

def create_cell_border():
    return Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

def export_inventory_to_excel(materials: List[Material]) -> io.BytesIO:
    """Xuất danh sách tồn kho nguyên vật liệu xưởng gỗ ra file Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ton_Kho_Vat_Tu"

    # Tiêu đề báo cáo
    ws.merge_cells("A1:G1")
    ws["A1"] = FACTORY_NAME.upper()
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"BÁO CÁO TỒN KHO NGUYÊN VẬT LIỆU (Ngày {date.today().strftime('%d/%m/%Y')})"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["STT", "Mã NVL", "Tên nguyên vật liệu", "Nhóm / Loại", "Đơn vị", "Tồn kho", "Đơn giá vốn (đ)", "Thành tiền (đ)"]
    ws.append([]) # Dòng 3 trống
    ws.append(headers) # Dòng 4

    header_fill, header_font, header_align, header_border = create_header_style()
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border

    row_num = 5
    border = create_cell_border()
    total_val = 0.0

    for idx, m in enumerate(materials, start=1):
        line_total = m.stock_quantity * m.unit_price
        total_val += line_total
        row_data = [
            idx,
            m.code,
            m.name,
            m.category,
            m.unit,
            m.stock_quantity,
            m.unit_price,
            line_total
        ]
        ws.append(row_data)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = border
            if col_idx in [6, 7, 8]:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in [1, 2, 5]:
                cell.alignment = Alignment(horizontal="center")
        row_num += 1

    # Dòng tổng cộng
    ws.merge_cells(f"A{row_num}:G{row_num}")
    total_label = ws.cell(row=row_num, column=1, value="TỔNG GIÁ TRỊ TỒN KHO")
    total_label.font = Font(name="Arial", size=11, bold=True)
    total_label.alignment = Alignment(horizontal="right")

    total_val_cell = ws.cell(row=row_num, column=8, value=total_val)
    total_val_cell.font = Font(name="Arial", size=11, bold=True, color="DC2626")
    total_val_cell.number_format = '#,##0'
    total_val_cell.alignment = Alignment(horizontal="right")

    # Điều chỉnh độ rộng cột
    col_widths = [8, 15, 35, 18, 12, 12, 18, 22]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_cashbook_to_excel(transactions: List[CashTransaction]) -> io.BytesIO:
    """Xuất sổ quỹ thu chi ra file Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "So_Quy_Thu_Chi"

    ws.merge_cells("A1:H1")
    ws["A1"] = FACTORY_NAME.upper()
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"SỔ QUỸ TIỀN MẶT & NGÂN HÀNG (Ngày xuất: {date.today().strftime('%d/%m/%Y')})"
    ws["A2"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["STT", "Mã phiếu", "Ngày", "Loại", "Danh mục", "Người nộp/nhận", "Phương thức", "Số tiền (đ)"]
    ws.append([])
    ws.append(headers)

    header_fill, header_font, header_align, header_border = create_header_style()
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border

    row_num = 5
    border = create_cell_border()
    total_in = 0.0
    total_out = 0.0

    for idx, tx in enumerate(transactions, start=1):
        if tx.transaction_type == "Thu":
            total_in += tx.amount
        else:
            total_out += tx.amount

        row_data = [
            idx,
            tx.code,
            tx.transaction_date.strftime("%d/%m/%Y"),
            tx.transaction_type,
            tx.category,
            tx.payer_or_receiver or "",
            tx.payment_method,
            tx.amount
        ]
        ws.append(row_data)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = border
            if col_idx == 8:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
                if tx.transaction_type == "Thu":
                    cell.font = Font(color="16A34A", bold=True)
                else:
                    cell.font = Font(color="DC2626", bold=True)
            elif col_idx in [1, 2, 3, 4, 7]:
                cell.alignment = Alignment(horizontal="center")
        row_num += 1

    # Điều chỉnh độ rộng cột
    col_widths = [8, 18, 14, 10, 25, 25, 16, 20]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
